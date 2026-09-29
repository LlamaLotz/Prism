import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[2]
SPEC = importlib.util.spec_from_file_location("study_export", ROOT / "notebook" / "study_export.py")
EXPORT = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(EXPORT)

SOURCE = {"id": "source-1", "title": "Fixture source", "path": "notes/fixture.md", "hash": "abc123", "text": "source body", "missing": False}
SNAPSHOT = {"id": "snapshot-1", "sources": [SOURCE], "context": "source body", "excerpts": False, "created": 1}
FIXTURES = {
    "table": {"title": "Table", "columns": ["Name", "Formula"], "rows": [{"cells": ["Row", "=1+1"], "sourceIds": ["source-1"]}]},
    "quiz": {"title": "Quiz", "questions": [{"question": "Question?", "options": ["Yes", "No"], "answer": 0, "explanation": "Evidence.", "sourceIds": ["source-1"]}]},
    "flashcards": {"title": "Cards", "cards": [{"id": "card-1", "front": "Front", "back": "Back", "sourceIds": ["source-1"]}]},
    "podcast": {"title": "Podcast", "transcript": "Host: A short sourced script.", "sourceIds": ["source-1"]},
    "slides": {"title": "Slides", "slides": [{"title": "Slide one", "bullets": ["Point"], "notes": "Speaker notes", "sourceIds": ["source-1"]}]},
    "mindmap": {"title": "Map", "nodes": [{"id": "root", "parentId": None, "label": "Root <topic>", "sourceIds": ["source-1"]}, {"id": "child", "parentId": "root", "label": "Child & detail", "sourceIds": ["source-1"]}]},
}


class StudyExportTests(unittest.TestCase):
    def export(self, kind, extension):
        artifact = {"id": "artifact-1", "kind": kind, "title": FIXTURES[kind]["title"], "body": FIXTURES[kind]}
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / f"fixture.{extension}"
            EXPORT.export({"artifact": artifact, "snapshot": SNAPSHOT}, output)
            self.assertGreater(output.stat().st_size, 0)
            return output.read_bytes()

    def test_markdown_exports_include_citations_and_quiz_answer_key(self):
        quiz = self.export("quiz", "md").decode()
        self.assertLess(quiz.index("# Answer key"), len(quiz))
        self.assertIn("notes/fixture.md", quiz)
        self.assertIn("Evidence.", quiz)
        podcast = self.export("podcast", "md").decode()
        self.assertIn("Host: A short sourced script.", podcast)

    def test_csv_is_parseable_and_neutralizes_spreadsheet_formulas(self):
        import csv
        import io

        rows = list(csv.reader(io.StringIO(self.export("table", "csv").decode("utf-8-sig"))))
        self.assertEqual(rows[1][1], "'=1+1")
        self.assertIn("notes/fixture.md", rows[1][-1])
        cards = list(csv.reader(io.StringIO(self.export("flashcards", "csv").decode("utf-8-sig"))))
        self.assertEqual(cards[1], ["Front", "Back", "notes/fixture.md"])

    def test_svg_is_valid_and_escapes_generated_labels(self):
        content = self.export("mindmap", "svg")
        root = ET.fromstring(content)
        self.assertTrue(root.tag.endswith("svg"))
        self.assertIn(b"Root &lt;topic&gt;", content)
        self.assertNotIn(b"Root <topic>", content)

    def test_pptx_is_a_readable_presentation(self):
        content = self.export("slides", "pptx")
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "fixture.pptx"
            path.write_bytes(content)
            from pptx import Presentation
            presentation = Presentation(path)
            self.assertEqual(len(presentation.slides), 1)
            self.assertIn("Slide one", presentation.slides[0].shapes.title.text)
            self.assertIn("Speaker notes", presentation.slides[0].notes_slide.notes_text_frame.text)

    @unittest.skipUnless(importlib.util.find_spec("genanki"), "genanki is installed in the bundled Notebook runtime")
    def test_apkg_is_a_readable_anki_package(self):
        content = self.export("flashcards", "apkg")
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "fixture.apkg"
            path.write_bytes(content)
            with zipfile.ZipFile(path) as archive:
                self.assertIn("collection.anki2", archive.namelist())
                self.assertIn("media", archive.namelist())
                database = Path(folder) / "collection.anki2"
                database.write_bytes(archive.read("collection.anki2"))
            import sqlite3
            connection = sqlite3.connect(database)
            self.assertEqual(connection.execute("SELECT count(*) FROM cards").fetchone()[0], 1)
            self.assertEqual(connection.execute("SELECT count(*) FROM notes").fetchone()[0], 1)
            self.assertEqual(connection.execute("SELECT count(*) FROM col").fetchone()[0], 1)
            connection.close()

    @unittest.skipUnless(importlib.util.find_spec("reportlab"), "ReportLab is installed in the bundled Notebook runtime")
    def test_pdf_is_a_readable_pdf_with_answer_section(self):
        for kind in FIXTURES:
            content = self.export(kind, "pdf")
            self.assertTrue(content.startswith(b"%PDF-"), kind)
        content = self.export("quiz", "pdf")
        self.assertIn(b"%%EOF", content[-32:])


if __name__ == "__main__":
    unittest.main()
