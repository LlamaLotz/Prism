"""Offline exporters for saved Prism artifacts. Never executes generated code or HTML."""
import csv
import hashlib
import html
import json
import re
import sys
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape


def sources(item, snapshot):
    ids = item.get('sourceIds', [])
    return '; '.join(s['path'] or s['title'] for s in snapshot['sources'] if s['id'] in ids)


def markdown(a, snap):
    b, kind = a['body'], a['kind']
    lines = ['# ' + a['title'], '']
    if b.get('shortfall'):
        lines += ['> ' + b['shortfall'], '']
    if kind == 'quiz':
        for i, q in enumerate(b['questions']):
            lines += [f"## {i + 1}. {q['question']}", *[f"- {chr(65+j)}. {o}" for j, o in enumerate(q['options'])], '']
        lines += ['# Answer key', '']
        for i, q in enumerate(b['questions']):
            lines += [f"{i + 1}. {chr(65+q['answer'])}: {q['explanation']}", 'Sources: ' + sources(q, snap), '']
    elif kind == 'flashcards':
        for card in b['cards']:
            lines += ['## ' + card['front'], card['back'], 'Sources: ' + sources(card, snap), '']
    elif kind == 'slides':
        for i, slide in enumerate(b['slides']):
            lines += [f"## {i+1}. {slide['title']}", *['- ' + x for x in slide['bullets']], 'Speaker notes: ' + slide['notes'], 'Sources: ' + sources(slide, snap), '']
    elif kind == 'podcast':
        lines += [b['transcript']]
    elif kind == 'mindmap':
        def visit(parent, depth):
            for node in b['nodes']:
                if node.get('parentId') == parent:
                    lines.append('  ' * depth + '- ' + node['label'] + ' — ' + sources(node, snap))
                    visit(node['id'], depth+1)
        visit(None, 0)
    elif kind == 'table':
        lines += [' | '.join(b['columns'])]
        for row in b['rows']:
            lines += [' | '.join('Unknown' if v is None else v for v in row['cells']), 'Sources: ' + sources(row, snap)]
    lines += ['', '## Source versions', *[f"- {s['path']}: {s['hash']}" for s in snap['sources']]]
    if snap.get('excerpts'):
        lines += ['', 'Generated from selected excerpts, not complete source coverage.']
    return '\n'.join(lines)

def anki_guid(artifact_id, card_id):
    """Generate the same stable, Anki-safe base91 GUID style as genanki."""
    alphabet = 'abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789!#$%&()*+,-./:;<=>?@[]^_`{|}~'
    digest = hashlib.sha256(f'{artifact_id}__{card_id}'.encode('utf-8')).digest()[:8]
    value = int.from_bytes(digest, 'big')
    chars = []
    while value:
        chars.append(alphabet[value % len(alphabet)])
        value //= len(alphabet)
    return ''.join(reversed(chars))


def export(data, output):
    a, snap = data['artifact'], data['snapshot']
    b, kind, ext = a['body'], a['kind'], output.suffix
    if ext == '.md':
        output.write_text(markdown(a, snap), encoding='utf-8')
    elif ext == '.csv':
        def safe(value):
            value = '' if value is None else str(value)
            return "'" + value if value.lstrip().startswith(('=', '+', '-', '@', '\t', '\r')) else value
        with output.open('w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            if kind == 'table':
                writer.writerow([safe(x) for x in b['columns']] + ['Sources'])
                for row in b['rows']:
                    writer.writerow([safe(x) for x in row['cells']] + [safe(sources(row, snap))])
            else:
                writer.writerow(['Front', 'Back', 'Sources'])
                for card in b['cards']:
                    writer.writerow([safe(card['front']), safe(card['back']), safe(sources(card, snap))])
    elif ext == '.pptx':
        from pptx import Presentation
        from pptx.util import Inches, Pt
        p = Presentation()
        p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
        for item in b['slides']:
            slide = p.slides.add_slide(p.slide_layouts[1])
            slide.shapes.title.text = item['title']
            frame = slide.placeholders[1].text_frame
            frame.clear()
            frame.word_wrap = True
            for i, text in enumerate(item['bullets']):
                paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                paragraph.text = text
                paragraph.font.size = Pt(22 if sum(map(len, item['bullets'])) < 650 else 16)
            slide.notes_slide.notes_text_frame.text = item['notes'] + '\nSources: ' + sources(item, snap)
        p.save(output)
    elif ext == '.apkg':
        import genanki
        model = genanki.Model(1895260421, 'Prism Basic', fields=[{'name': n} for n in ['Front', 'Back', 'Sources']], templates=[{'name': 'Card 1', 'qfmt': '{{Front}}', 'afmt': '{{FrontSide}}<hr id="answer">{{Back}}<br><small>{{Sources}}</small>'}])
        deck_id = int(hashlib.sha256(a['id'].encode()).hexdigest()[:12], 16)
        deck = genanki.Deck(deck_id, a['title'])
        collection_ids = {}
        import sqlite3
        import tempfile
        import time
        timestamp = int(time.time())
        def stable_id(value):
            return int.from_bytes(hashlib.sha256(value.encode('utf-8')).digest()[:7], 'big') + 1
        deck.deck_id = stable_id('prism-deck:' + a['id'])
        for model_field in model.fields:
            model_field['id'] = stable_id('prism-field:' + model_field['name'])
        for index, card in enumerate(b['cards']):
            note = genanki.Note(model=model, fields=[html.escape(card['front']), html.escape(card['back']), html.escape(sources(card, snap))], guid=anki_guid(a['id'], card['id']))
            note.due = index + 1
            deck.add_note(note)
        with tempfile.TemporaryDirectory(prefix='prism-apkg-') as folder:
            package = Path(folder) / 'package.apkg'
            genanki.Package(deck).write_to_file(str(package), timestamp=timestamp)
            with zipfile.ZipFile(package) as archive:
                database = Path(folder) / 'collection.anki2'
                database.write_bytes(archive.read('collection.anki2'))
            db = sqlite3.connect(database)
            row = db.execute('SELECT decks FROM col').fetchone()
            decks = json.loads(row[0])
            deck_key = str(deck.deck_id)
            decks[deck_key]['id'] = int.from_bytes(hashlib.sha256(a['id'].encode('utf-8')).digest()[:6], 'big') + 1
            decks[deck_key]['name'] = a['title']
            decks[deck_key]['mod'] = timestamp
            db.execute('UPDATE col SET decks = ?', (json.dumps(decks),))
            db.commit()
            db.close()
            media = json.dumps({'0': 'empty'})
            with zipfile.ZipFile(output, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
                archive.write(database, 'collection.anki2')
                archive.writestr('media', media)

    elif ext == '.svg':
        nodes = b['nodes']
        positions = {}
        by_id = {node['id']: node for node in nodes}
        child_rows = {}
        for node in nodes:
            child_rows.setdefault(node.get('parentId'), []).append(node)
        def depth(node):
            return 0 if node.get('parentId') is None else 1 + depth(by_id[node['parentId']])
        def assign(parent, row):
            for node in child_rows.get(parent, []):
                positions[node['id']] = (30 + depth(node) * 320, 30 + row[0] * 70)
                row[0] += 1
                assign(node['id'], row)
        assign(None, [0])
        height = len(nodes)*70+40
        parts = [f'<svg xmlns="http://www.w3.org/2000/svg" width="1000" height="{height}" viewBox="0 0 1000 {height}">', '<rect width="100%" height="100%" fill="white"/>']
        for node in nodes:
            x,y = positions[node['id']]
            if node.get('parentId'):
                px,py = positions[node['parentId']]
                parts.append(f'<path d="M {px+260} {py+20} L {x} {y+20}" fill="none" stroke="#475569"/>')
            parts += [f'<g><title>{html.escape(node["label"] + " — " + sources(node,snap))}</title><rect x="{x}" y="{y}" width="280" height="45" rx="8" fill="#e0f2fe" stroke="#0369a1"/>', f'<text x="{x+10}" y="{y+27}" font-family="sans-serif" font-size="13">{html.escape(node["label"][:38])}</text></g>']
        output.write_text(''.join(parts)+'</svg>', encoding='utf-8')
    elif ext == '.pdf':
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        import reportlab
        font = Path(reportlab.__file__).parent / 'fonts' / 'Vera.ttf'
        bold = Path(reportlab.__file__).parent / 'fonts' / 'VeraBd.ttf'
        pdfmetrics.registerFont(TTFont('Prism', str(font)))
        pdfmetrics.registerFont(TTFont('Prism-Bold', str(bold)))
        styles = getSampleStyleSheet()
        for style in styles.byName.values(): style.fontName = 'Prism'
        styles['Heading1'].fontName = 'Prism-Bold'
        styles['Heading2'].fontName = 'Prism-Bold'
        styles['BodyText'].leading = 15
        story = []
        for line in markdown(a,snap).splitlines():
            if line == '# Answer key': story.append(PageBreak())
            if line.startswith('# '): style,text=styles['Heading1'],line[2:]
            elif line.startswith('## '): style,text=styles['Heading2'],line[3:]
            elif line.startswith('- '): style,text=styles['BodyText'],u'•  '+line[2:]
            elif line.startswith('> '): style,text=styles['Italic'],line[2:]
            else: style,text=styles['BodyText'],line
            story.append(Paragraph(escape(text), style) if text else Spacer(1, 8))
        SimpleDocTemplate(str(output), title=a['title']).build(story)
    else:
        raise ValueError('Unsupported export format')


if __name__ == '__main__':
    export(json.loads(Path(sys.argv[1]).read_text(encoding='utf-8')), Path(sys.argv[2]))
